import zipfile

import pytest
from openpyxl import Workbook

from app.imports.parsers import ImportFileError, detect_format, parse_file
from app.services.import_service import ImportService


LIMITS = {
    "maximum_cell_length": 100,
    "maximum_columns": 20,
    "maximum_rows": 100,
    "maximum_worksheets": 5,
}


def _docx(path, rows):
    table_rows = "".join(
        "<w:tr>{}</w:tr>".format(
            "".join(
                "<w:tc><w:p><w:r><w:t>{}</w:t></w:r></w:p></w:tc>".format(value)
                for value in row
            )
        )
        for row in rows
    )
    document = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body><w:tbl>{table_rows}</w:tbl></w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", document)


@pytest.mark.parametrize(
    ("filename", "content_type", "sample", "expected"),
    [
        ("inventory.csv", "text/csv", b"hostname,ip\nhost1,10.0.0.1", "csv"),
        ("inventory.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", b"PK\x03\x04...", "xlsx"),
        ("inventory.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", b"PK\x03\x04...", "docx"),
    ],
)
def test_supported_format_detection(filename, content_type, sample, expected):
    assert detect_format(filename, content_type, sample) == expected


def test_docx_table_headers_and_inventory_rows_are_extracted(tmp_path):
    path = tmp_path / "inventory.docx"
    _docx(path, [["Hostname", "IP Address"], ["host-1", "10.0.0.1"]])

    parsed = parse_file(path, "docx", LIMITS)

    assert parsed.headers == ["Hostname", "IP Address"]
    assert parsed.rows == [(2, {"Hostname": "host-1", "IP Address": "10.0.0.1"}, [])]


def test_docx_preserves_rows_with_inconsistent_column_counts(tmp_path):
    path = tmp_path / "inventory.docx"
    _docx(path, [["Hostname", "IP Address"], ["host-1"]])

    parsed = parse_file(path, "docx", LIMITS)

    assert parsed.rows[0][1] == {"Hostname": "host-1", "IP Address": None}
    assert parsed.rows[0][2] == [
        {
            "field": "row",
            "code": "column_count",
            "message": "Row column count differs from the header",
        }
    ]


def test_docx_without_inventory_table_is_rejected(tmp_path):
    path = tmp_path / "empty.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")

    with pytest.raises(ImportFileError, match="must contain a table"):
        parse_file(path, "docx", LIMITS)


def test_xlsx_skips_title_row_and_ignores_trailing_blank_columns(tmp_path):
    path = tmp_path / "desktop-allocation.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Allocation"
    sheet.merge_cells("A1:G1")
    sheet["A1"] = "Desktop Allocation - assigned inventory"
    for index, header in enumerate(
        ["SN", "Laptop Description & Config", "Item Serial Number", "System Name", "Staff Assigned", "Department", "Date of Deployment", None, None],
        start=1,
    ):
        sheet.cell(row=2, column=index, value=header)
    for index, value in enumerate([1, "Laptop config", "SN-001", "HOST-01", "Alex", "IT", "2024-03-05"], start=1):
        sheet.cell(row=3, column=index, value=value)
    workbook.save(path)

    parsed = parse_file(path, "xlsx", LIMITS)

    assert parsed.headers == [
        "SN", "Laptop Description & Config", "Item Serial Number", "System Name",
        "Staff Assigned", "Department", "Date of Deployment",
    ]
    assert parsed.rows[0][0] == 3
    assert parsed.rows[0][1]["System Name"] == "HOST-01"
    assert parsed.rows[0][1]["Date of Deployment"] == "2024-03-05"
    assert parsed.rows[0][2] == []


def test_csv_uses_recognized_header_after_title_row_and_keeps_source_row_numbers(tmp_path):
    path = tmp_path / "inventory.csv"
    path.write_text(
        "Allocation summary,,,,\nSN,Item Serial Number,System Name,Department,,\n1,SN-001,HOST-01,IT,,\n",
        encoding="utf-8",
    )

    parsed = parse_file(path, "csv", LIMITS)

    assert parsed.headers == ["SN", "Item Serial Number", "System Name", "Department"]
    assert parsed.rows[0][0] == 3
    assert parsed.rows[0][1]["System Name"] == "HOST-01"


def test_repeated_hostname_is_reported_as_a_duplicate_identifier():
    keys = ImportService._duplicate_key({"hostname": "host-01"})

    assert ("hostname", "host-01") in keys
